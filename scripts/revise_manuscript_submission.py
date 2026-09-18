#!/usr/bin/env python3
"""Create the analysis-audited manuscript revision from an accepted DOCX copy."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


def clear_paragraph(paragraph: Paragraph) -> None:
	for child in list(paragraph._p):
		if child.tag != qn("w:pPr"):
			paragraph._p.remove(child)


def set_paragraph(paragraph: Paragraph, text: str) -> None:
	clear_paragraph(paragraph)
	run = paragraph.add_run(text)
	run.font.size = Pt(11)


def find_one(doc: Document, prefix: str) -> Paragraph:
	matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
	if len(matches) != 1:
		raise RuntimeError(
			f"Expected one paragraph beginning {prefix!r}; found {len(matches)}"
		)
	return matches[0]


def delete_paragraph(paragraph: Paragraph) -> None:
	parent = paragraph._element.getparent()
	parent.remove(paragraph._element)
	paragraph._p = paragraph._element = None


def delete_between(doc: Document, start_prefix: str, end_prefix: str) -> None:
	paragraphs = doc.paragraphs
	starts = [
		i for i, paragraph in enumerate(paragraphs)
		if paragraph.text.strip().startswith(start_prefix)
	]
	ends = [
		i for i, paragraph in enumerate(paragraphs)
		if paragraph.text.strip().startswith(end_prefix)
	]
	if len(starts) != 1 or len(ends) != 1 or starts[0] >= ends[0]:
		raise RuntimeError(
			f"Could not identify one ordered block from {start_prefix!r} "
			f"to {end_prefix!r}"
		)
	for paragraph in reversed(paragraphs[starts[0]:ends[0]]):
		delete_paragraph(paragraph)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
	new_p = OxmlElement("w:p")
	paragraph._p.addnext(new_p)
	return Paragraph(new_p, paragraph._parent)


def preceding_paragraph(paragraph: Paragraph) -> Paragraph:
	previous = paragraph._p.getprevious()
	if previous is None or previous.tag != qn("w:p"):
		raise RuntimeError("Expected an image paragraph immediately before caption")
	return Paragraph(previous, paragraph._parent)


def set_picture(paragraph: Paragraph, path: Path, width_inches: float) -> None:
	clear_paragraph(paragraph)
	paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
	paragraph.paragraph_format.keep_with_next = True
	paragraph.paragraph_format.space_after = Pt(3)
	paragraph.add_run().add_picture(str(path), width=Inches(width_inches))


def set_caption(paragraph: Paragraph, lead: str, detail: str) -> None:
	clear_paragraph(paragraph)
	paragraph.style = "normal"
	paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
	paragraph.paragraph_format.keep_together = True
	paragraph.paragraph_format.space_before = Pt(3)
	paragraph.paragraph_format.space_after = Pt(8)
	lead_run = paragraph.add_run(lead)
	lead_run.bold = True
	lead_run.font.size = Pt(10)
	detail_run = paragraph.add_run(" " + detail)
	detail_run.font.size = Pt(10)


def replace_text(doc: Document, replacements: dict[str, str]) -> None:
	for prefix, text in replacements.items():
		set_paragraph(find_one(doc, prefix), text)


def remove_page_break(paragraph: Paragraph) -> None:
	for line_break in paragraph._p.xpath('.//w:br[@w:type="page"]'):
		line_break.getparent().remove(line_break)


def normalize_revision_artifacts(doc: Document) -> None:
	"""Remove residual review colours, highlighting, shading, and border bars."""
	for run_properties in doc.element.body.xpath(".//w:rPr"):
		for tag in ("w:color", "w:highlight", "w:shd", "w:bdr"):
			for element in list(run_properties.findall(qn(tag))):
				run_properties.remove(element)
	for paragraph_properties in doc.element.body.xpath(".//w:pPr"):
		for tag in ("w:shd", "w:pBdr"):
			for element in list(paragraph_properties.findall(qn(tag))):
				paragraph_properties.remove(element)
	# Accepting inserted/deleted text does not necessarily remove recorded
	# formatting changes. Those records cause revision bars and can make older
	# colours reappear in Word/LibreOffice, so remove the change records while
	# keeping the current formatting state.
	change_tags = (
		"w:rPrChange",
		"w:pPrChange",
		"w:sectPrChange",
		"w:tblPrChange",
		"w:trPrChange",
		"w:tcPrChange",
		"w:numPrChange",
	)
	for tag in change_tags:
		for element in list(doc.element.body.xpath(f".//{tag}")):
			parent = element.getparent()
			if parent is not None:
				parent.remove(element)


def main() -> None:
	parser = argparse.ArgumentParser()
	parser.add_argument("--input", required=True, type=Path)
	parser.add_argument("--output", required=True, type=Path)
	parser.add_argument(
		"--figure-dir",
		type=Path,
		default=Path("figures/main"),
	)
	args = parser.parse_args()

	doc = Document(str(args.input))

	# Submission copies must not contain the internal journal, reviewer, or task notes
	# that precede the abstract in the working draft.
	delete_between(doc, "Potential journals (ladder):", "Abstract")

	replacements = {
		"Classical theory posits forest succession": (
			"Successional theory predicts that forest communities become more "
			"predictable as stands mature, but whether this pattern persists across "
			"environmental contexts remains unclear. Using 49,377 inventory "
			"observations from 27,270 permanent plots spanning United States forests, "
			"we modelled community-weighted mean expression of nine functional traits "
			"as a function of stand age and abiotic context. Environmental predictors "
			"collectively received more SHAP attribution than stand age in every "
			"trait-forest combination (median environmental-to-age ratios 1.7-31.1; "
			"all empirical intervals from PID-grouped cross-validation exceeded one), "
			"although this comparison aggregates five environmental axes against one "
			"age variable. Environmental context also modified modelled successional "
			"rates in 61 of 90 trait-gradient-forest combinations, with stronger rate "
			"modification in broadleaf forests. These responses did not follow a "
			"universal convergence, divergence, or trait-syndrome pattern: environmental "
			"trait separation widened in 39 combinations, narrowed in 20, and remained "
			"uncertain in 31. PID-grouped cross-validation further showed that "
			"predictive skill was higher in later than early succession, while "
			"environmental differences in predictive skill persisted in most "
			"late-successional cells. Forest functional succession is therefore "
			"environmentally conditional, but the environment does not impose a single "
			"direction of change. Forecasts of forest recovery should represent stand "
			"age and environmental context jointly and retain context-specific uncertainty."
		),
		"Here we address these questions using data from over 49,000 permanent forest plots": (
			"Here we address these questions using 49,377 inventory observations from "
			"27,270 permanent forest plots spanning the United States and Alaska, drawn "
			"from the US Forest Inventory and Analysis programme. Using plot-level "
			"measurements of tree size and composition, we modelled community-weighted "
			"mean expression of nine functional traits spanning leaf economics "
			"(specific leaf area, leaf potassium, leaf density), hydraulics (conduit "
			"diameter), bark structure (bark thickness), whole-tree architecture (tree "
			"height), resource acquisition (root depth), reproduction (seed dry mass), "
			"and light competition (shade tolerance) as a function of stand age and "
			"abiotic predictors. We then compared environmental and age-related model "
			"attribution, quantified how environmental conditions modify modelled "
			"successional trait trajectories, and assessed how trait predictability "
			"changes through succession across contrasting environmental contexts."
		),
		"We propose three hypotheses": (
			"We tested three predictions. First, environmental predictors should receive "
			"greater collective model attribution than stand age, with the weakest "
			"environmental dominance for traits most directly related to light "
			"competition and whole-tree architecture (Grime 1977; Kraft et al. 2015). "
			"Second, environmental context should modify successional trait trajectories; "
			"we specifically predicted faster change in resource-acquisitive traits under "
			"productive conditions and the converse for stress-tolerant traits (Boukili "
			"& Chazdon 2017; Lohbeck et al. 2014; Walker & Chapin 1987). Third, trait "
			"predictability should be higher in later succession but remain contingent "
			"on environmental context, such that predictability differences between "
			"contrasting environments persist rather than converge toward zero (Chase "
			"2010; HilleRisLambers et al. 2012; Poorter et al. 2023)."
		),
		"Environmental Dominance of Functional Trait Variation": (
			"Environmental Dominance of Model Attribution"
		),
		"Predictability Divergence Through Succession": (
			"Environmental Contingency of Trait Predictability"
		),
		"We used data from the US Forest Inventory and Analysis": (
			"We used data from the US Forest Inventory and Analysis (FIA) programme "
			"(Gray et al. 2012), a long-term monitoring dataset recording tree "
			"occurrences across standardised permanent plots throughout the United "
			"States. From the full FIA database, we retained living trees identified to "
			"species level, measured under standard fixed-area plot designs (design "
			"codes 1, 501, and 505), recorded in a single homogeneous condition class, "
			"and inventoried from 1980 onwards. We excluded privately owned plots and "
			"plots under active forestry management. Where plots had been measured more "
			"than once, we retained up to three of the most recent inventories per plot. "
			"Stand age is reported by FIA as the average total age of live trees within "
			"the plot's stand-size class (Gray et al. 2012) and was used as a proxy for "
			"successional stage. We excluded observations in the upper 10% of the stand "
			"age distribution to focus on approximately 1-150 years and reduce the "
			"influence of sparsely sampled old-growth stands. Broadleaf forests comprised "
			"temperate broadleaf and Mediterranean woodland biomes; coniferous forests "
			"comprised boreal and temperate conifer biomes. The final dataset contained "
			"49,377 inventory observations from 27,270 permanent plots. Analyses were "
			"conducted separately for broadleaf and coniferous forests."
		),
		"We modelled CWM trait expression as a function of stand age": (
			"We modelled CWM trait expression as a function of stand age and five abiotic "
			"predictors using random forest regression implemented in ranger (Wright & "
			"Ziegler 2017) in R v.4.5.0. Random forests accommodate nonlinear "
			"relationships and predictor interactions without prespecifying their form "
			"(Breiman 2001). Separate models were fitted for each trait and forest type "
			"(18 models). Hyperparameters (number of trees, variables per split, and "
			"minimum node size) were selected by out-of-bag grid search and used "
			"consistently in the grouped resampling analyses below. Final-model held-out "
			"R² and RMSE are reported as descriptive diagnostics; formal uncertainty "
			"and predictive-skill analyses kept all inventories from the same permanent "
			"plot together within resampling partitions."
		),
		"To quantify the relative contribution of environmental and successional filtering": (
			"To compare environmental and age-related model attribution, we computed "
			"SHAP values using fastshap (Greenwell 2024). For each trait and forest type, "
			"we summed absolute SHAP values for five environmental predictors "
			"(temperature, soil water retention, precipitation, elevation, and soil pH) "
			"and divided this sum by the absolute SHAP attribution of stand age. This "
			"ratio compares the combined attribution of five environmental axes with one "
			"age variable and should not be interpreted as a per-predictor contrast. To "
			"quantify partition stability, we repeated 10-fold cross-validation five "
			"times while assigning every inventory from the same permanent plot (PID) to "
			"the same fold. Models were fitted to nine folds and SHAP values were "
			"estimated for up to 1,500 observations in the held-out fold using 50 Monte "
			"Carlo simulations. Figure 2a shows the median ratio and empirical 2.5th-97.5th "
			"percentile interval across the 50 partitions. Figure 2b descriptively "
			"decomposes total absolute SHAP attribution across individual predictors in "
			"the final models."
		),
		"To test whether environmental conditions modulate successional trait trajectories": (
			"To test whether environmental conditions modify modelled successional trait "
			"trajectories, we used a PID-cluster bootstrap. Within each forest type, "
			"complete permanent plots were sampled with replacement, retaining every "
			"inventory from each selected plot. For each of the five environmental "
			"predictors in turn, the resampled data were divided into lower and upper "
			"quartiles. Separate random forest models were fitted within each stratum "
			"using stand age and all five abiotic predictors, with the tuned "
			"hyperparameters from the corresponding global model. Each trait-forest "
			"combination was repeated for 100 bootstrap draws."
		),
		"For each stratified model, we quantified the successional trait trajectory": (
			"For each stratified model, we estimated partial dependence on a common "
			"stand-age grid from 10 to 100 years in five-year increments while retaining "
			"the observed values of all other predictors. Every environmental stratum "
			"contained at least 50 inventory observations, at least 25 resampled PID "
			"clusters, and observed age support across the full grid. The resulting "
			"bootstrap trajectories are shown in Fig. S-7."
		),
		"We first asked whether environmental context altered the rate and direction": (
			"We summarised average directional change by fitting a linear trend to each "
			"partial-dependence trajectory. The signed contrast was Δslope = "
			"slope upper - slope lower; positive values indicate a more positive "
			"trajectory in the upper environmental quartile, not necessarily a faster "
			"absolute rate. We therefore used |Δslope| to quantify the magnitude of "
			"environmental rate modification and compared |slope upper| - |slope lower| "
			"when asking which stratum changed faster. Empirical 95% PID-cluster "
			"bootstrap intervals were the 2.5th and 97.5th percentiles across 100 draws."
		),
		"We next asked whether contrasting environmental contexts converged or diverged": (
			"We next asked whether environmental separation in modelled trait expression "
			"changed through succession. Within each bootstrap draw, we calculated the "
			"absolute upper-lower difference in predicted trait expression at stand ages "
			"10 and 100. A positive paired change indicates widening environmental "
			"separation and a negative change indicates narrowing; change was classified "
			"as supported when its 95% bootstrap interval excluded zero."
		),
		"where MSE is the mean squared error": (
			"where the numerator is the out-of-fold sum of squared prediction errors and "
			"the denominator is the observed sum of squares around the mean. VEcv is a "
			"scale-free out-of-fold analogue of R². We used 30 repeats of 10-fold "
			"cross-validation within each forest type, assigning all inventories from the "
			"same PID to the same fold. Within each repeat, out-of-fold predictions were "
			"classified into lower and upper environmental quartiles and 10-year stand-age "
			"bins. Cells with fewer than 30 observations were excluded. Medians and "
			"empirical 95% repeated-partition intervals were calculated across repeats; "
			"these intervals quantify sensitivity to data partitioning rather than "
			"population sampling uncertainty. Legates-McCabe efficiency (E1), based on "
			"absolute rather than squared errors, was used as a robustness check."
		),
		"To quantify divergence in predictability across environmental gradients": (
			"For each trait, environmental gradient, age bin, and repeat, we calculated "
			"ΔVEcv = VEcv upper - VEcv lower and its absolute magnitude. Signed "
			"contrasts identify which stratum was more predictable; absolute contrasts "
			"measure the strength of environmental dependence without cancellation among "
			"traits. Figure 4 uses the common, well-supported age window of 15-125 years. "
			"Predictive skill was averaged with equal weight across traits after averaging "
			"their ten environmental strata. Environmental predictability gaps were "
			"paired within trait and repeat, converted to absolute differences, and then "
			"averaged across traits for each gradient."
		),
		"VEcv =": (
			"VEcv = 1 - sum[(observed - predicted)^2] / "
			"sum[(observed - mean observed)^2]"
		),
		"Random forest models explained a mean of 73.2%": (
			"Final-model held-out diagnostics explained a mean of 73.2% of trait variance "
			"(R² range 0.576-0.890) across traits and forest types. Because these original "
			"train-test splits were not grouped by PID, we treat them as descriptive model "
			"diagnostics rather than inferential estimates; grouped out-of-fold predictive "
			"skill is reported below."
		),
		"Environmental predictors collectively explained substantially more trait variation": (
			"The combined SHAP attribution of the five environmental predictors exceeded "
			"that of stand age for all nine traits and both forest types (Fig. 2a). "
			"PID-grouped cross-validation medians ranged from 1.69 for coniferous tree "
			"height to 31.1 for coniferous seed dry mass, and every empirical 95% "
			"partition interval remained above one (minimum lower bound 1.58). Mean "
			"environment-to-age ratios were 10.1 in broadleaf and 20.5 in coniferous "
			"forests (medians 10.9 and 25.7). Individual-predictor decomposition remained "
			"heterogeneous (Fig. 2b): temperature contributed the largest share for all "
			"nine broadleaf traits, whereas coniferous traits were divided among "
			"temperature, elevation, and soil pH. Because the numerator combines five "
			"predictors, these results support strong collective environmental attribution "
			"rather than the dominance of every individual environmental variable."
		),
		"Stand age interacted with environmental conditions to modulate the rate": (
			"Environmental context commonly modified modelled successional trajectories "
			"(Fig. 3; Fig. S-7). Empirical PID-cluster bootstrap intervals for the signed "
			"slope difference excluded zero in 34 of 45 broadleaf and 27 of 45 coniferous "
			"trait-gradient combinations. Median slope-modification magnitudes were larger "
			"in broadleaf forests for every environmental gradient, ranging from 0.154 to "
			"0.711 trait SD per 100 years, compared with 0.047-0.133 in coniferous forests "
			"(Fig. 3b). However, the stratum with the faster absolute rate was not "
			"consistent: upper environmental quartiles changed faster in 17 combinations, "
			"lower quartiles in 29, and 44 were unresolved. Environmental separation in "
			"modelled trait expression likewise changed in both directions between ages "
			"10 and 100 (Fig. 3a): 39 combinations widened, 20 narrowed, and 31 were "
			"uncertain. No environmental gradient produced a universal response across "
			"traits and forest types."
		),
		"Trait predictability increased through succession in both forest types": (
			"Trait predictive skill was higher in later than early succession in both "
			"forest types, although trajectories were non-monotonic (Fig. 4a; Fig. S-8). "
			"Across repeated PID-grouped partitions, mean VEcv increased from 0.171 at "
			"early ages to 0.346 at late ages in broadleaf forests (paired change 0.174, "
			"empirical 95% interval 0.170-0.183) and from 0.366 to 0.456 in coniferous "
			"forests (change 0.091, interval 0.087-0.094). Environmental dependence of "
			"predictive skill persisted through later succession (Fig. 4b): mean absolute "
			"upper-lower gaps across traits ranged from 0.090 to 0.243 among the ten "
			"gradient-forest combinations at late ages. Signed repeated-partition intervals "
			"excluded zero in 86.7% of broadleaf and 92.6% of coniferous "
			"trait-gradient-age cells at ages 105-125. Gap trajectories strengthened for "
			"some gradients and weakened for others; grouped resampling therefore did not "
			"support a general narrowing pattern or a unique temperature exception."
		),
		"Forest functional communities are widely assumed": (
			"Forest functional communities are expected to become more predictable as "
			"succession proceeds, yet their trajectories unfold within spatially variable "
			"environmental constraints. Across United States forests, environmental "
			"predictors collectively received substantially more model attribution than "
			"stand age, environmental context commonly modified modelled successional "
			"rates, and environmental differences in predictive skill persisted into "
			"later succession. At the same time, trajectory responses were directionally "
			"heterogeneous rather than organised around one universal convergence, "
			"divergence, or trait-syndrome pattern. The strongest supported synthesis is "
			"therefore not that the environment fixes a single alternative successional "
			"pathway, but that functional succession and its predictability remain "
			"conditional on environmental context."
		),
		"The near-universal dominance of environmental over successional predictors": (
			"The consistent environmental-to-age SHAP ratios extend evidence that climate "
			"and soils strongly structure forest functional composition (Laughlin et al. "
			"2012; Maynard et al. 2022; Swenson & Enquist 2007). Tree height and shade "
			"tolerance had the lowest ratios, consistent with a comparatively large role "
			"for stand development and light competition. Ratios were generally larger in "
			"coniferous than broadleaf forests, but this descriptive contrast should not "
			"be assigned a mechanism from SHAP attribution alone. Moreover, the ratio "
			"compares five environmental predictors with one age variable. Its robust "
			"departure from one demonstrates collective environmental attribution, not "
			"that each environmental predictor individually dominates stand age or that "
			"the attributed relationships are causal."
		),
		"The widespread robust differences in successional slopes": (
			"The grouped bootstrap confirms that environmental context frequently changes "
			"modelled successional trajectories, especially in broadleaf forests. However, "
			"the direction of signed slope contrasts did not translate into a general "
			"difference in absolute rate between productive and stressful environments. "
			"This distinction matters because a more positive slope can still represent "
			"slower change when both trajectories are negative. The predicted universal "
			"acquisitive-versus-conservative response was therefore not supported across "
			"traits and gradients. Rather than weakening the main conclusion, this result "
			"shows that environmental modulation is widespread but functionally "
			"multidimensional, with its direction depending on the trait, gradient, and "
			"forest type."
		),
		"The strikingly weak modulation of successional slopes by temperature": (
			"Changes in environmental trait separation reinforce this heterogeneous view. "
			"Widening was more frequent than narrowing overall, but both outcomes were "
			"common and nearly one third of combinations remained uncertain. Broadleaf "
			"temperature and coniferous precipitation showed the most consistent widening, "
			"whereas coniferous soil water retention showed no supported widening. These "
			"patterns reject a single convergence rule while cautioning against replacing "
			"it with an equally simple divergence rule. Environmental context changes the "
			"functional distance between trajectories, but the direction of that change is "
			"not universal."
		),
		"The finding that trait predictability increased through succession": (
			"The higher predictive skill in later than early succession is consistent with "
			"increasing regularity as stands mature (Chase 2010; Fukami 2015; "
			"HilleRisLambers et al. 2012), although the non-monotonic age profiles do not "
			"support a continuous rise. Crucially, environmental predictability gaps "
			"remained substantial at late ages. Their trajectories differed among "
			"gradients and forest types rather than following a general narrowing pattern "
			"or a temperature-specific exception. We therefore interpret Figure 4 as "
			"evidence for persistent context dependence, "
			"not for a temperature-specific mechanism. VEcv measures out-of-fold predictive "
			"skill conditional on the measured predictors; it is informative about "
			"predictability but is not, by itself, a direct measure of ecological "
			"determinism."
		),
		"In conclusion, we show that the environmental context does not merely modulate succession": (
			"Several limitations constrain interpretation. Stand age is inferred from the "
			"age of dominant trees and may covary with site productivity and the measured "
			"environmental predictors, potentially inflating apparent environmental-age "
			"contrasts. The cross-sectional design means that modelled trajectories "
			"compare different stands across ages rather than directly observing within-plot "
			"succession. PID-grouped cross-validation and cluster bootstrap prevent repeated "
			"inventories from crossing resampling partitions, but they do not remove "
			"confounding by disturbance history, species pools, land-use legacies, or "
			"unmeasured site conditions. Finally, community-weighted means describe central "
			"functional composition but omit functional diversity, intraspecific variation, "
			"and ontogenetic trait change."
		),
		"The cross-sectional design means": (
			"These limitations make causal language inappropriate, but they do not erase "
			"the convergent evidence across three grouped analyses: environmental context "
			"dominates collective model attribution, commonly modifies modelled age "
			"relationships, and remains associated with predictive skill in later "
			"succession. Future longitudinal analyses can test whether the same patterns "
			"emerge from within-plot change and whether disturbance history explains the "
			"remaining heterogeneity among trait trajectories."
		),
		"These findings contribute to our understanding how functional recovery": (
			"These findings have direct implications for forecasting forest recovery. "
			"Models that represent succession only as time since disturbance, with climate "
			"as a secondary correction to a regionally fixed trajectory, will miss "
			"environment-dependent differences in functional change and predictive "
			"uncertainty (Cook-Patton et al. 2020; Poorter et al. 2021; Pugh et al. 2019). "
			"Our results support treating environmental context as a primary axis alongside "
			"stand age while allowing response direction to vary among traits and forest "
			"types. This formulation is deliberately mechanism-neutral and provides a "
			"robust basis for evaluating recovery forecasts under changing climates."
		),
	}

	replace_text(doc, replacements)

	# Remove the superseded PDP window-sensitivity and projected quantile-sensitivity
	# paragraphs; neither is part of the grouped production analyses.
	for prefix in (
		"(|Delta|0)",
		"(|Δ|₀)",
		"Finally, we tested whether conclusions about environmental differences",
		"(0-150 years)",
		"(0–150 years)",
	):
		matches = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
		for paragraph in matches:
			delete_paragraph(paragraph)

	figure2 = args.figure_dir / "fig2_shap_with_cv_uncertainty.png"
	figure3 = args.figure_dir / "fig3_pdp_grouped_review.png"
	figure4 = args.figure_dir / "fig4_vecv_grouped_review.png"
	for path in (figure2, figure3, figure4):
		if not path.exists():
			raise FileNotFoundError(path)

	fig2_caption = find_one(doc, "Figure 2.")
	set_picture(preceding_paragraph(fig2_caption), figure2, 6.27)
	set_caption(
		fig2_caption,
		"Figure 2. Environmental context dominates model attribution of functional trait variation.",
		"(a) Ratio of the summed absolute SHAP attribution of five environmental predictors to the absolute attribution of stand age (log₂ scale). Points are medians across five repeats of PID-grouped 10-fold cross-validation; horizontal lines are empirical 95% repeated-partition intervals. The dashed line at 1× denotes equal attribution. Traits are ordered from greatest to lowest environmental dominance. (b) Descriptive final-model decomposition of total absolute SHAP attribution among individual predictors. Stand age (brown) is shown at the top of each bar.",
	)

	fig3_old = find_one(doc, "Figure 3.")
	set_picture(fig3_old, figure3, 6.27)
	fig3_caption = insert_paragraph_after(fig3_old)
	set_caption(
		fig3_caption,
		"Figure 3. Environmental context modifies successional trait trajectories without imposing a universal direction.",
		"(a) Absolute difference in modelled trait expression between upper and lower environmental quartiles at stand ages 10 and 100. Points are trait-gradient combinations; shapes identify gradients and colours classify the paired change as widening, narrowing, or uncertain according to the 95% PID-cluster bootstrap interval. Horizontal and vertical bars show bootstrap intervals for early and late separation. (b) Magnitude of the signed slope difference between environmental quartiles, expressed in trait SD per 100 years. Points are traits, faded points have intervals including zero, diamonds are gradient medians, and labels give the number of supported signed contrasts out of nine traits.",
	)

	fig4_caption = find_one(doc, "Figure 4.")
	set_picture(preceding_paragraph(fig4_caption), figure4, 6.27)
	set_caption(
		fig4_caption,
		"Figure 4. Trait predictive skill is higher later in succession but remains environmentally contingent.",
		"(a) Out-of-fold variance explained (VEcv), averaged with equal weight across nine traits after averaging their ten environmental strata. Points and lines are medians across 30 repeats of PID-grouped 10-fold cross-validation; ribbons are empirical 95% repeated-partition intervals. (b) Mean absolute upper-lower environmental difference in VEcv across traits for each gradient. Zero indicates identical predictive skill between environmental quartiles. Both panels use the common supported age-bin midpoint range of 15-125 years.",
	)

	# Keep the existing Figure 1 caption intact on one page.
	find_one(doc, "Figure 1.").paragraph_format.keep_together = True

	# Removing internal notes and superseded methods changes pagination. The old
	# manual breaks before Introduction and Results then create blank or mostly
	# blank pages, so allow those sections to flow naturally.
	for heading_text in ("Introduction", "Results"):
		heading = find_one(doc, heading_text)
		previous = heading._p.getprevious()
		if previous is not None and previous.tag == qn("w:p"):
			remove_page_break(Paragraph(previous, heading._parent))

	normalize_revision_artifacts(doc)

	args.output.parent.mkdir(parents=True, exist_ok=True)
	doc.save(str(args.output))
	print(f"Wrote {args.output}")


if __name__ == "__main__":
	main()
